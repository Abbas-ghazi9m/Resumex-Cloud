import io
from typing import Dict, Any, List
from docx import Document

class DOCXParsingError(Exception):
    pass

def extract_docx_data(docx_bytes: bytes) -> Dict[str, Any]:
    """
    Parses DOCX document using python-docx.
    Extracts paragraphs, headings, bullet points, and tables.
    """
    if not docx_bytes or len(docx_bytes) < 32:
        raise DOCXParsingError("File is empty or corrupted (too small for valid DOCX).")
    
    # Check PK zip header (DOCX is a zip archive)
    if not docx_bytes.startswith(b"PK\x03\x04"):
        raise DOCXParsingError("Invalid DOCX format: missing standard OpenXML package header.")

    try:
        doc = Document(io.BytesIO(docx_bytes))
    except Exception as e:
        raise DOCXParsingError(f"Corrupted or unreadable DOCX archive: {str(e)}")

    paragraphs_text = []
    headings = []
    bullet_points = []
    has_tables = len(doc.tables) > 0

    for p in doc.paragraphs:
        txt = p.text.strip()
        if not txt:
            continue
        paragraphs_text.append(txt)
        
        # Check style
        style_name = p.style.name.lower() if p.style else ""
        if "heading" in style_name or "title" in style_name:
            headings.append(txt)
        elif "list" in style_name or txt.startswith(("-", "•", "*", "–")):
            bullet_points.append(txt)

    # Also parse tables if any
    table_texts = []
    for table in doc.tables:
        for row in table.rows:
            row_cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if row_cells:
                table_texts.append(" | ".join(row_cells))

    combined = "\n".join(paragraphs_text)
    if table_texts:
        combined += "\n\n" + "\n".join(table_texts)

    combined = combined.strip()
    word_count = len(combined.split())

    if word_count == 0:
        raise DOCXParsingError("DOCX document contains no text content.")

    # Estimate page count (standard approx 350-450 words per page in resumes)
    page_count = max(1, round(word_count / 380))

    return {
        "text": combined,
        "page_count": page_count,
        "word_count": word_count,
        "headings": headings,
        "bullet_count": len(bullet_points),
        "has_tables": has_tables,
        "table_count": len(doc.tables),
    }
